from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
heading = doc.add_heading('Rule-Based Recipe Matching and Ranking Algorithm', level=1)

# Input/Output
doc.add_paragraph('Input: D (Set of ingredients detected by the YOLOv8 model)', style='List Bullet')
doc.add_paragraph('Output: R_top5 (List of top 5 recommended recipes)', style='List Bullet')

doc.add_paragraph('\n1. INITIALIZATION:')
p1 = doc.add_paragraph()
p1.add_run('    Let I_db be the set of all ingredients in the database\n')
p1.add_run('    Let R_db be the set of all recipes in the database\n')
p1.add_run('    Let M be an empty list for matched recipes\n')
p1.add_run('    Let D_ids be an empty set for detected ingredient IDs')

doc.add_paragraph('\n2. INGREDIENT MAPPING:')
p2 = doc.add_paragraph()
p2.add_run('    FOR EACH detected item d IN D DO\n').bold = True
p2.add_run('        FOR EACH ingredient i IN I_db DO\n').bold = True
p2.add_run('            IF d matches i.name OR d matches i.coco_class THEN\n').bold = True
p2.add_run('                Add i.id to D_ids\n')
p2.add_run('            END IF\n').bold = True
p2.add_run('        END FOR\n').bold = True
p2.add_run('    END FOR').bold = True

doc.add_paragraph('\n3. RULE-BASED SCORING:')
p3 = doc.add_paragraph()
p3.add_run('    FOR EACH recipe r IN R_db DO\n').bold = True
p3.add_run('        Let Req be the set of required ingredient IDs for r\n')
p3.add_run('        Let match_score = 0\n')
p3.add_run('        FOR EACH req_id IN Req DO\n').bold = True
p3.add_run('            IF req_id is present in D_ids THEN\n').bold = True
p3.add_run('                match_score = match_score + 1\n')
p3.add_run('            END IF\n').bold = True
p3.add_run('        END FOR\n').bold = True
p3.add_run('        IF match_score > 0 THEN\n').bold = True
p3.add_run('            Set r.score = match_score\n')
p3.add_run('            Set r.total_req = length(Req)\n')
p3.add_run('            Append r to M\n')
p3.add_run('        END IF\n').bold = True
p3.add_run('    END FOR').bold = True

doc.add_paragraph('\n4. RANKING AND FILTERING:')
p4 = doc.add_paragraph()
p4.add_run('    Sort M in descending order based on r.score\n')
p4.add_run('    IF two recipes have the same r.score THEN\n').bold = True
p4.add_run('        Sort them in ascending order based on r.total_req (Tie-breaker rule)\n')
p4.add_run('    END IF\n\n').bold = True
p4.add_run('    // Restrict output to top 5 to reduce cognitive load on mobile UI\n').italic = True
p4.add_run('    Let R_top5 = M[0 ... 4]')

doc.add_paragraph('\n5. RETURN R_top5').bold = True

doc.save(r'C:\Users\Marck\Downloads\Algorithm_Pseudocode.docx')
print("Document created successfully.")
